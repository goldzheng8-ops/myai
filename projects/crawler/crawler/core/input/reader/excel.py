from collections.abc import Iterable
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

from core.input.reader.base import BaseDataReader
from core.input.record import DataRecord
from core.input.source import DataFormat, DataSource


class ExcelDataReader(
    BaseDataReader,
):
    format = DataFormat.EXCEL

    def _read(
        self,
        source: DataSource,
    ) -> Iterable[DataRecord]:
        path = Path(source.path)

        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        try:
            worksheet = workbook.active

            if worksheet is None:
                raise ValueError(
                    f"Excel workbook in {path} "
                    "does not contain an active worksheet.",
                )

            rows = worksheet.iter_rows(
                values_only=True,
            )

            headers = next(
                rows,
                None,
            )

            if headers is None:
                raise ValueError(
                    f"Excel worksheet in {path} "
                    "is empty.",
                )

            normalized_headers = tuple(
                self._normalize_header(
                    value,
                    index,
                )
                for index, value in enumerate(
                    headers,
                )
            )

            for row in rows:
                values = {
                    header: value
                    for header, value in zip(
                        normalized_headers,
                        row,
                    )
                }

                yield DataRecord(
                    values=values,
                )

        finally:
            workbook.close()

    def _normalize_header(
        self,
        value: Any,
        index: int,
    ) -> str:
        if value is None:
            raise ValueError(
                f"Excel header at column "
                f"{index + 1} is empty.",
            )

        return str(value)